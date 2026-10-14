"""Gera o modelo de posts da Mary Criativa para importar no Canva.

Saida: Modelo-Posts-Mary-Criativa.pptx (1080 x 1350, formato retrato do Instagram).
Uso: python3 gerar_modelo.py
"""
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

AQUI = Path(__file__).parent
W, H = 1080, 1350
PX = 9525  # EMU por pixel

# Paleta da Mary Criativa
ROSA = "E07A9C"
ROSA_ESCURO = "C95A80"
ROSA_CLARO = "FBE4EC"
LILAS = "B08AD6"
LILAS_ESCURO = "8A63B5"
LILAS_CLARO = "F1E8FA"
DOURADO = "C49A48"
DOURADO_CLARO = "F3E3BC"
TINTA = "4A3550"
TINTA_SUAVE = "7C6A82"
CREME = "FFFBF8"
BRANCO = "FFFFFF"

SCRIPT = "Great Vibes"
SERIF = "Cormorant Garamond"
SANS = "Montserrat"


def rgb(h):
    return RGBColor.from_string(h)


# ---------- fundos (imagens) ----------
def fundo_xadrez(nome, cor_listra=(176, 138, 214), base=(253, 250, 255), passo=54, alpha=40):
    img = Image.new("RGB", (W, H), base)
    over = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(over)
    for x in range(0, W, passo * 2):
        d.rectangle([x, 0, x + passo - 1, H], fill=cor_listra + (alpha,))
    for y in range(0, H, passo * 2):
        d.rectangle([0, y, W, y + passo - 1], fill=cor_listra + (alpha,))
    img = Image.alpha_composite(img.convert("RGBA"), over).convert("RGB")
    caminho = AQUI / nome
    img.save(caminho)
    return caminho


# ---------- ajudantes ----------
def rect(slide, x, y, w, h, fill=None, line=None, line_w=2, dash=False, forma=MSO_SHAPE.RECTANGLE, raio=None, rot=0, nome=None):
    s = slide.shapes.add_shape(forma, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    if raio is not None and forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = raio
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w * 0.75)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.rotation = rot
    if nome:
        s.name = nome
    s.text_frame.text = ""
    return s


def texto(slide, x, y, w, h, linhas, fonte=SANS, tam=24, cor=TINTA, negrito=False, italico=False,
          alinhar=PP_ALIGN.CENTER, vert=MSO_ANCHOR.MIDDLE, espaco=1.0, letras=0, rot=0, riscado=False, nome=None):
    tb = slide.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tb.rotation = rot
    if nome:
        tb.name = nome
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = vert
    if isinstance(linhas, str):
        linhas = [linhas]
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhar
        p.line_spacing = espaco
        r = p.add_run()
        r.text = linha
        f = r.font
        f.name = fonte
        f.size = Pt(tam * 0.75)
        f.bold = negrito
        f.italic = italico
        f.color.rgb = rgb(cor)
        rpr = r._r.get_or_add_rPr()
        if letras:
            rpr.set("spc", str(int(letras * 75)))
        if riscado:
            rpr.set("strike", "sngStrike")
    return tb


def espaco_arte(slide, x, y, w, h, rotulo, forma=MSO_SHAPE.ROUNDED_RECTANGLE, raio=0.08, rot=0, cor=LILAS, fundo="F8F3FD"):
    """Retangulo tracejado para colocar a arte real (apague depois)."""
    s = rect(slide, x, y, w, h, fill=fundo, line=cor, line_w=2.5, dash=True, forma=forma, raio=raio, rot=rot, nome="Espaco para arte: " + rotulo)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_top = Emu(int(min(h, 400) * 0.05 * PX))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "SUA ARTE AQUI"
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = SANS, Pt(max(9, min(w, h) * 0.075) * 0.75), True, rgb(cor)
    r._r.get_or_add_rPr().set("spc", "150")
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = rotulo
    r2.font.name, r2.font.size, r2.font.italic, r2.font.color.rgb = SERIF, Pt(max(11, min(w, h) * 0.07) * 0.75), True, rgb(TINTA_SUAVE)
    return s


def selo(slide, x, y, d=150):
    """Selo redondo de criacao autoral, como o do post de referencia."""
    rect(slide, x, y, d, d, fill=BRANCO, line=ROSA, line_w=2, dash=True, forma=MSO_SHAPE.OVAL, nome="Selo criacao autoral")
    texto(slide, x, y + d * 0.16, d, d * 0.14, "CRIAÇÃO AUTORAL", SANS, d * 0.065, LILAS_ESCURO, negrito=True, letras=1)
    texto(slide, x, y + d * 0.32, d, d * 0.3, "Mary", SCRIPT, d * 0.3, ROSA_ESCURO)
    texto(slide, x, y + d * 0.58, d, d * 0.12, "CRIATIVA", SANS, d * 0.06, DOURADO, negrito=True, letras=3)
    texto(slide, x, y + d * 0.72, d, d * 0.12, "@marycriativa", SERIF, d * 0.075, TINTA_SUAVE, italico=True)


def fita(slide, cx, cy, w, h, txt, cor_fita=LILAS_CLARO, cor_txt=LILAS_ESCURO, borda=LILAS):
    """Faixa (fita) com texto, ex.: '10 anos'."""
    pontas = 0.16
    for lado in (-1, 1):
        px = cx + lado * (w / 2 - w * pontas * 0.4) - w * pontas / 2
        rect(slide, px, cy - h / 2 + h * 0.28, w * pontas, h, fill=cor_fita, line=borda, line_w=1.5,
             forma=MSO_SHAPE.CHEVRON if lado == 1 else MSO_SHAPE.CHEVRON, rot=0 if lado == 1 else 180)
    rect(slide, cx - w * 0.4, cy - h / 2, w * 0.8, h, fill=cor_fita, line=borda, line_w=1.5, nome="Fita")
    texto(slide, cx - w * 0.4, cy - h / 2, w * 0.8, h, txt, SCRIPT, h * 0.62, cor_txt)


def rodape(slide, txt="Desenvolvida por @marycriativa", cor=LILAS_ESCURO, y=H - 78):
    texto(slide, 0, y, W, 40, txt, SERIF, 30, cor, italico=True, letras=1)


def flores_texto(slide, x, y, txt="✦", cor=DOURADO, tam=34):
    texto(slide, x, y, tam * 1.2, tam * 1.2, txt, SANS, tam, cor)


# ---------- slides ----------
prs = Presentation()
prs.slide_width = Emu(W * PX)
prs.slide_height = Emu(H * PX)
BRANCO_LAYOUT = prs.slide_layouts[6]

xadrez = fundo_xadrez("fundo-xadrez-lilas.png")
xadrez_rosa = fundo_xadrez("fundo-xadrez-rosa.png", cor_listra=(224, 122, 156), base=(255, 250, 252), alpha=46)


def novo(fundo=None, cor=None):
    s = prs.slides.add_slide(BRANCO_LAYOUT)
    if fundo:
        pic = s.shapes.add_picture(str(fundo), 0, 0, prs.slide_width, prs.slide_height)
        pic.name = "Fundo"
    elif cor:
        rect(s, 0, 0, W, H, fill=cor, nome="Fundo")
    return s


# 1. Capa do projeto: brasao no centro
s = novo(xadrez)
espaco_arte(s, 150, 190, 780, 860, "brasão do projeto (PNG sem fundo)", raio=0.06, fundo="FFFFFF")
texto(s, 150, 470, 780, 260, "Luna", SCRIPT, 230, DOURADO, nome="Nome")
fita(s, W / 2, 900, 560, 92, "10 anos")
selo(s, 870, 1130)
texto(s, 60, 1130, 700, 50, "Identidade visual · Tema Jardim Encantado", SERIF, 30, LILAS_ESCURO, italico=True, alinhar=PP_ALIGN.LEFT)

# 2. Monograma, logotipo e paleta
s = novo(cor=BRANCO)
cores = [DOURADO, "F2D16B", "FBE9C6", "EBD9F0", "CDB6DC", LILAS]
faixa_y, faixa_h = 100, 1000
for i, c in enumerate(cores):
    rect(s, 960, faixa_y + i * (faixa_h / 6) + 6, 120, faixa_h / 6 - 12, fill=c, nome=f"Cor {i + 1}")
espaco_arte(s, 120, 90, 520, 500, "monograma (letra inicial)", raio=0.06)
texto(s, 120, 90, 520, 500, "L", SERIF, 420, "CDB6DC", negrito=True, nome="Monograma")
espaco_arte(s, 40, 680, 880, 420, "logotipo do projeto", raio=0.06)
texto(s, 40, 690, 880, 300, "Luna", SCRIPT, 250, LILAS, nome="Logotipo")
fita(s, 480, 1000, 460, 80, "10 anos", cor_fita="FBEFD0", cor_txt=DOURADO, borda="E6CF94")
selo(s, 690, 470, 140)
rodape(s)

# 3. Estampas e elementos
s = novo(cor=CREME)
texto(s, 0, 70, W, 70, "ESTAMPAS EXCLUSIVAS", SANS, 30, LILAS_ESCURO, negrito=True, letras=6)
rect(s, W / 2 - 50, 150, 100, 3, fill=DOURADO)
for i in range(4):
    cx, cy = 110 + (i % 2) * 450, 200 + (i // 2) * 450
    espaco_arte(s, cx, cy, 410, 410, f"estampa {i + 1}", raio=0.04)
texto(s, 0, 1100, W, 70, "Elementos", SCRIPT, 80, ROSA_ESCURO)
for i in range(5):
    espaco_arte(s, 110 + i * 178, 1180, 150, 110, f"elemento {i + 1}", raio=0.2, cor=ROSA, fundo="FFF3F7")

# 4. Aplicacoes: convite, save the date e placa
s = novo(xadrez_rosa)
texto(s, 0, 60, W, 110, "Aplicações", SCRIPT, 110, ROSA_ESCURO)
texto(s, 0, 170, W, 40, "CONVITE · SAVE THE DATE · PLACA DE BOAS-VINDAS", SANS, 22, LILAS_ESCURO, negrito=True, letras=3)
for i, (x, y, rot, rot_txt) in enumerate([(90, 300, -6, "convite"), (380, 360, 5, "save the date")]):
    rect(s, x, y, 330, 620, fill="2E2433", forma=MSO_SHAPE.ROUNDED_RECTANGLE, raio=0.1, rot=rot, nome=f"Celular {i + 1}")
    espaco_arte(s, x + 16, y + 30, 298, 560, rot_txt, raio=0.08, rot=rot, fundo="FFFFFF")
rect(s, 760, 1000, 14, 300, fill="C89B6D", rot=12, nome="Cavalete")
rect(s, 940, 1000, 14, 300, fill="C89B6D", rot=-12, nome="Cavalete")
rect(s, 850, 1020, 14, 260, fill="B88A5C", nome="Cavalete")
espaco_arte(s, 720, 420, 280, 600, "placa de boas-vindas", raio=0.04, fundo="FFFFFF")
selo(s, 80, 1110, 140)

# 5. Chamada final
s = novo(xadrez)
rect(s, 110, 260, 860, 830, fill=BRANCO, line=DOURADO, line_w=2, forma=MSO_SHAPE.ROUNDED_RECTANGLE, raio=0.06, nome="Cartao")
texto(s, 160, 330, 760, 110, "Quer uma identidade", SERIF, 70, TINTA, negrito=True)
texto(s, 160, 430, 760, 180, "assim?", SCRIPT, 170, ROSA_ESCURO)
texto(s, 180, 640, 720, 130, ["Cada festa merece uma identidade só dela.", "Me chama que eu crio a da sua!"], SERIF, 36, TINTA_SUAVE, italico=True, espaco=1.1)
rect(s, 290, 820, 500, 96, fill=ROSA, forma=MSO_SHAPE.ROUNDED_RECTANGLE, raio=0.5, nome="Botao")
texto(s, 290, 820, 500, 96, "ORÇAMENTO NO WHATSAPP", SANS, 26, BRANCO, negrito=True, letras=2)
texto(s, 160, 950, 760, 60, "@marycriativa", SERIF, 40, LILAS_ESCURO, italico=True)
texto(s, 0, 1150, W, 120, "Mary Criativa", SCRIPT, 90, DOURADO)

# 6. Post de promocao de pacote
s = novo(cor=BRANCO)
texto(s, 0, 60, W, 100, "PROMOÇÃO", SANS, 84, ROSA_ESCURO, negrito=True, letras=4, nome="Titulo")
texto(s, 0, 160, W, 40, "Festa infantil  •  Pacote de identidade visual", SANS, 24, LILAS_ESCURO)
for i, c in enumerate([ROSA, LILAS, DOURADO, ROSA, LILAS, DOURADO]):
    rect(s, 330 + i * 72, 222, 52, 4, fill=c)
texto(s, 80, 270, 420, 36, "PACOTE 02", SANS, 24, LILAS_ESCURO, negrito=True, alinhar=PP_ALIGN.LEFT, letras=2)
texto(s, 80, 305, 440, 110, "Encanto", SCRIPT, 120, ROSA_ESCURO, alinhar=PP_ALIGN.LEFT, nome="Nome do pacote")
rect(s, 80, 425, 400, 3, fill=DOURADO)
texto(s, 80, 445, 420, 40, "de R$ 650,00", SANS, 30, TINTA_SUAVE, alinhar=PP_ALIGN.LEFT, riscado=True, nome="Preco antigo")
texto(s, 80, 485, 420, 40, "por apenas", SANS, 30, TINTA_SUAVE, alinhar=PP_ALIGN.LEFT)
texto(s, 80, 525, 440, 100, "R$ 560,00", SANS, 84, ROSA_ESCURO, negrito=True, alinhar=PP_ALIGN.LEFT, nome="Preco")
rect(s, 545, 270, 3, 350, fill=LILAS)
texto(s, 580, 270, 440, 40, "ITENS INCLUSOS", SANS, 26, LILAS_ESCURO, negrito=True, alinhar=PP_ALIGN.LEFT, letras=2)
col1 = ["Brasão", "Logotipo", "Monogramas", "Paleta de cores", "3 estampas", "Elementos"]
col2 = ["Tipografia", "Convite virtual", "Save the date", "Cardápio", "2 tags"]
texto(s, 580, 325, 230, 290, ["•  " + t for t in col1], SANS, 23, TINTA, alinhar=PP_ALIGN.LEFT, vert=MSO_ANCHOR.TOP, espaco=1.35, nome="Itens 1")
texto(s, 810, 325, 230, 290, ["•  " + t for t in col2], SANS, 23, TINTA, alinhar=PP_ALIGN.LEFT, vert=MSO_ANCHOR.TOP, espaco=1.35, nome="Itens 2")
# colagem
espaco_arte(s, 60, 680, 340, 250, "brasão", raio=0.3)
espaco_arte(s, 420, 690, 260, 160, "logotipo", raio=0.1)
espaco_arte(s, 720, 670, 260, 230, "monograma", raio=0.1)
for i, c in enumerate([ROSA, "F7C6D6", LILAS, DOURADO, "FBF3E8", "E9DBF7"]):
    rect(s, 440 + (i % 3) * 82, 880 + (i // 3) * 82, 70, 70, fill=c, line="EADDF2" if c in ("FBF3E8",) else None, line_w=1, forma=MSO_SHAPE.OVAL, nome=f"Cor {i + 1}")
espaco_arte(s, 60, 960, 180, 180, "estampa 1", raio=0.04)
espaco_arte(s, 200, 1080, 180, 180, "estampa 2", raio=0.04)
espaco_arte(s, 60, 1140, 130, 130, "estampa 3", raio=0.04)
for i, (x, y, rot) in enumerate([(440, 1030, -10), (575, 1060, 6)]):
    rect(s, x, y, 150, 195, fill="2E2433", forma=MSO_SHAPE.ROUNDED_RECTANGLE, raio=0.1, rot=rot, nome=f"Tablet {i + 1}")
    espaco_arte(s, x + 11, y + 11, 128, 173, "convite" if i else "save the date", raio=0.06, rot=rot, fundo="FFFFFF")
espaco_arte(s, 790, 940, 200, 270, "placa", raio=0.04, fundo="FFFFFF")
rect(s, 820, 1200, 10, 110, fill="C89B6D", rot=10)
rect(s, 950, 1200, 10, 110, fill="C89B6D", rot=-10)
texto(s, 0, 1280, W, 60, "Mary Criativa", SCRIPT, 56, DOURADO)

saida = AQUI / "Modelo-Posts-Mary-Criativa.pptx"
prs.save(saida)
print("ok", saida)
